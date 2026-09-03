import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_proposal_docx(filename):
    doc = docx.Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styling colors
    DARK_BLUE = RGBColor(11, 25, 44)      # Primary Heading Color
    CYAN_ACCENT = RGBColor(0, 150, 214)   # Secondary Accent
    TEXT_DARK = RGBColor(44, 53, 64)     # Body Text
    WHITE = RGBColor(255, 255, 255)

    # Default Style Font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = TEXT_DARK

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(24)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24)
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.color.rgb = CYAN_ACCENT
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = CYAN_ACCENT
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = DARK_BLUE
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Segoe UI'
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_BLUE
        r_text = p.add_run(text)
        r_text.font.name = 'Segoe UI'
        r_text.font.color.rgb = TEXT_DARK
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Segoe UI'
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_BLUE
        r_text = p.add_run(text)
        r_text.font.name = 'Segoe UI'
        r_text.font.color.rgb = TEXT_DARK
        return p

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    # --- COVER / TITLE HEADER ---
    add_title("AMASAMYA: Inclusive Technology & Professional Learning Mega-Platform")
    add_subtitle("A Strategic Proposal & Master Execution Blueprint for Equal Opportunity in Tech, Language, and Career Empowerment\nCo-Authored for Review & Inputs by L. Subramani")

    # Metadata Block
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(18)
    r_meta = p_meta.add_run("Project Leaders: Akhilesh Malani (Accessibility Architect & Engineer) & L. Subramani (Journalist, Author & Communication Strategist)\nDate: August 2026 | Version: 1.0 (Draft Proposal)")
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 110, 120)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 1: EXECUTIVE VISION ---
    add_h1("1. Executive Vision & Core Mission")
    add_p("The mission of the AMASAMYA Mega-Platform is to create a unified, barrier-free ecosystem where persons with disabilities and mainstream society learn, build, and work hand-in-hand without any gaps. Accessibility must not be confined to software code audits or legal compliance checklists; it must extend to total human empowerment, language mastery, professional communication, and career independence.")
    
    add_p("AMASAMYA brings together two formidable pillars of leadership:")
    add_bullet(" Creator of the AMASAMYA accessibility audit suite across Android (Build 11), Chrome/Edge/Firefox extensions (v5.3), Web Audit Platform, VPAT 2.4 generator, and Indian national framework engines (GIGW 3.0 and IS 17802).", bold_prefix="Akhilesh Malani (Technology & Accessibility Lead):")
    add_bullet(" Veteran journalist with 28+ years of experience (Chief Copy Editor at Deccan Herald), acclaimed author of 'Lights Out', communications strategist, and disability advocate who has mastered executive English communication as a visually impaired leader.", bold_prefix="L. Subramani (English Language & Professional Communication Lead):")

    add_p("Together, this platform ensures that a learner - whether visually impaired, motor impaired, neurodivergent, or a non-native English speaker - receives world-class technical skills and executive-level written and verbal English mastery under one trusted brand.")

    # --- SECTION 2: MARKET NEED & GAP ANALYSIS ---
    add_h1("2. Market Landscape & Strategic Gap Analysis")
    add_p("An extensive study of Indian and global solutions reveals critical gaps that prevent disabled individuals and non-native learners from reaching full career independence:")

    # Table of Gap Analysis
    table_gap = doc.add_table(rows=1, cols=3)
    table_gap.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_gap.rows[0].cells
    hdr_titles = ["Market Sector", "Current Available Solutions & Limitations", "The AMASAMYA Advantage"]
    for i, title in enumerate(hdr_titles):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0B192C")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.size = Pt(10)

    gaps_data = [
        ("Indian Government EdTech (DIKSHA, SWAYAM, Sugamya)", "High course volume, but heavily reliant on inaccessible PDFs, image-only quizzes, and uncaptioned video streams. Sugamya offers accessible books but lacks interactive tech/code labs.", "Provides screen-reader native interactive learning, accessible code sandboxes, and full GIGW 3.0 / IS 17802 compliance."),
        ("Global Accessibility Training (Deque Univ, W3C WAI)", "Expensive ($150 - $300 USD/course), theoretical, desktop-centric, English-only, and lacks coverage of Indian national standards.", "Affordable/Free tier for learners, interactive 'Learn-by-Auditing' sandboxes, and multi-lingual regional language support."),
        ("English Learning Apps (Duolingo, Grammarly, ELSA)", "Heavy visual drag-and-drop mechanics unusable by screen reader users; focuses on casual phrases rather than corporate/journalism writing.", "Screen-reader first English language pedagogy curated by L. Subramani, covering basic grammar up to executive corporate writing."),
        ("Global LMS Platforms (Canvas, Moodle, Coursera)", "Passive file checkers that point out errors without teaching remediation; authoring tools lock out visually impaired educators.", "Accessible LMS + Course Builder allowing blind educators to author lessons, record audio-first modules, and manage courses.")
    ]

    for sector, limit, adv in gaps_data:
        row_cells = table_gap.add_row().cells
        row_cells[0].text = sector
        row_cells[1].text = limit
        row_cells[2].text = adv
        for idx, cell in enumerate(row_cells):
            set_cell_background(cell, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.color.rgb = TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 3: PLATFORM ARCHITECTURE ---
    add_h1("3. Core Pillars of the AMASAMYA Mega-Platform")
    add_p("The mega-platform is structured into four interconnected operational pillars:")

    add_h2("Pillar 1: School of English & Professional Communication (Led by L. Subramani)")
    add_p("Curated specifically for screen reader users, braille display readers, and non-native speakers seeking corporate fluency:")
    add_bullet(" Phonics, sentence structure, screen-reader friendly grammar drills, and listening comprehension.", bold_prefix="Track A - English Foundations:")
    add_bullet(" High-impact business email writing, technical documentation, copyediting, and eliminating ambiguity.", bold_prefix="Track B - Professional & Corporate Writing:")
    add_bullet(" Public speaking, meeting leadership, voice modulation, and interview mastery.", bold_prefix="Track C - Executive Verbal Fluency:")
    add_bullet(" Feature writing, blogging, alt-text authoring, and digital advocacy storytelling.", bold_prefix="Track D - Journalism & Content Craft:")

    add_h2("Pillar 2: Accessibility & Software Engineering Ecosystem (Led by Akhilesh Malani)")
    add_p("Leverages the existing AMASAMYA software suite:")
    add_bullet(" Auditing UIs against WCAG 2.2 AA, GIGW 3.0, and IS 17802 across Android and Web.", bold_prefix="Multi-Platform Engine:")
    add_bullet(" Browser sandbox allowing developers to write live code and receive instant TalkBack / screen reader utterance simulations.", bold_prefix="AMASAMYA CodeLab:")
    add_bullet(" One-click generation of official compliance documents for government and enterprise deployment.", bold_prefix="VPAT 2.4 ACR Exporter:")

    add_h2("Pillar 3: AMASAMYA CopyAudit & Writing Assistant")
    add_p("A specialized writing evaluation engine built into the platform:")
    add_bullet(" Evaluates student essays and emails for grammatical accuracy, tone, jargon density, and clarity.", bold_prefix="Automated Readability & Grammar Evaluation:")
    add_bullet(" Speech rate and pronunciation feedback tuned for screen reader users and non-native speakers.", bold_prefix="Audio & Speech Cadence Analysis:")

    add_h2("Pillar 4: Blind-First Accessible LMS & Authoring Engine")
    add_p("Empowers educators with disabilities to teach:")
    add_bullet(" Keyboard-only, high-contrast, screen-reader optimized student portal.", bold_prefix="Accessible Learner Experience:")
    add_bullet(" Accessible content creation tool allowing blind teachers to record audio micro-lectures, design quizzes, and publish courses without visual drag-and-drop dependencies.", bold_prefix="Inclusive Educator Dashboard:")

    # --- SECTION 4: CURRICULUM BLUEPRINT FOR L. SUBRAMANI ---
    add_h1("4. Proposed English & Communication Curriculum Blueprint")
    add_p("This draft curriculum structure is submitted specifically for review, refinement, and expansion by L. Subramani:")

    # Curriculum Table
    table_curr = doc.add_table(rows=1, cols=3)
    table_curr.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdr = table_curr.rows[0].cells
    c_hdr[0].text = "Module Level"
    c_hdr[1].text = "Core Competencies & Topics"
    c_hdr[2].text = "Pedagogical Approach & Audio Tools"
    for cell in c_hdr:
        set_cell_background(cell, "0096D6")
        p = cell.paragraphs[0]
        for r in p.runs:
            r.font.bold = True
            r.font.color.rgb = WHITE
            r.font.size = Pt(10)

    curr_data = [
        ("Level 1: Foundation English", "Parts of Speech, Tense Consistency, Sentence Construction, Vocabulary Expansion, Phonetics.", "Audio-first grammar drills, spoken word repetition labs, braille display layout exercises."),
        ("Level 2: Corporate Communication", "Business Emails, Status Reports, Technical Summaries, Copyediting, Removing Jargon.", "AMASAMYA CopyAudit engine feedback, real-world case study rewrites, proofreading exercises."),
        ("Level 3: Executive Speech & Interviews", "Interview Preparation, Pitching Ideas, Public Speaking, Voice Modulation, Managing Meetings.", "Recorded voice submission labs, mock interview practice, screen reader speed modulation tips."),
        ("Level 4: Masterclass in Writing & Journalism", "Long-form Feature Writing, Authoring Books/Articles, Digital Advocacy, Accessible Alt-Text Writing.", "Mentorship masterclasses by L. Subramani, peer review circles, publication opportunities.")
    ]

    for lvl, comp, ped in curr_data:
        row_cells = table_curr.add_row().cells
        row_cells[0].text = lvl
        row_cells[1].text = comp
        row_cells[2].text = ped
        for cell in row_cells:
            set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.color.rgb = TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- SECTION 5: IMPLEMENTATION ROADMAP ---
    add_h1("5. Phased 12-Month Implementation Roadmap")
    add_p("To ensure structured and sustainable execution, the mega-project is organized into three 4-month phases:")

    add_bullet(" Finalize curriculum outline with L. Subramani; launch initial accessible LMS web portal featuring English Foundation and WCAG/GIGW 3.0 introductory modules; integrate 1-click 'Learn in Academy' links inside AMASAMYA browser extensions and Android app.", bold_prefix="Phase 1 - Foundation & Ecosystem Integration (Months 1 to 4):")
    add_bullet(" Deploy AMASAMYA CodeLab (interactive code sandbox) and AMASAMYA CopyAudit (writing feedback engine); release audio-first spoken language practice labs; launch dual certification tracks.", bold_prefix="Phase 2 - Interactive Sandboxes & Writing Engines (Months 5 to 8):")
    add_bullet(" Release accessible course authoring dashboard for blind educators; establish institutional partnerships with universities, blind schools, and corporate DEI initiatives; launch Indic regional language learning paths.", bold_prefix="Phase 3 - Educator Authoring & Institutional Scaling (Months 9 to 12):")

    # --- SECTION 6: FEEDBACK REQUESTS FOR L. SUBRAMANI ---
    add_h1("6. Key Discussion Points & Feedback Requests for L. Subramani")
    add_p("Akhilesh Malani invites L. Subramani to review this strategic proposal and share insights on the following key areas:")
    add_bullet(" Does the proposed 4-level English curriculum capture the most essential writing and verbal skills needed for career independence?", bold_prefix="1. Curriculum Structure & Scope:")
    add_bullet(" How can we best design audio-first exercises that help screen reader users transition from computer audio to natural, persuasive human speaking skills?", bold_prefix="2. Audio-First Teaching Methodology:")
    add_bullet(" What specific features should the AMASAMYA CopyAudit writing evaluation tool prioritize to help students learn copyediting and clarity?", bold_prefix="3. CopyAudit Tool Integration:")
    add_bullet(" What approaches will best encourage blind students, non-native speakers, and professionals to engage in live masterclasses and peer writing circles?", bold_prefix="4. Mentorship & Student Engagement:")

    doc.save(filename)
    print(f"Proposal successfully created at: {filename}")

if __name__ == "__main__":
    create_proposal_docx("C:\\Users\\akhi_\\antigravity\\focused-fermi\\amasamya_mega_platform_proposal.docx")
